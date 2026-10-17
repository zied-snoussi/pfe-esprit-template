"""Builds the Word version of Cahier-des-charges.md (python-docx, no template needed).

Markdown subset: front matter (read for the header/footer, then skipped), # title, ## subtitle (first one)
then sections, ### subsections, tables, "- " bullets ("- [ ]" becomes a checkbox), **bold** and `code` inline.
The line <!-- diagrammes --> inserts every PNG of ../diagrammes/png/ with a caption built from its file name
(01-cas-utilisation-global.png -> "Diagramme 01 : Cas utilisation global"); override it in CAPTIONS.

Usage: python build_cahier_docx.py [source.md] [output.docx]
Then:  soffice --headless --convert-to pdf <output.docx>   (or open it in Word, right-click the table of
       contents -> Update field, then File -> Export -> PDF)
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for landscape tweaks)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Cahier-des-charges.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "Cahier-des-charges.docx"
PNG = ROOT / "diagrammes" / "png"
# Optional caption per diagram number; by default the caption comes from the file name.
CAPTIONS = {
    "01": "Cas d'utilisation global", "02": "Diagramme de classes d'analyse",
    "03": "Séquence : prise de rendez-vous", "04": "Activité : prise de rendez-vous",
    "05": "États : rendez-vous", "06": "Déploiement",
}
META = {}  # front matter of the source (projet, version, ...), filled by convert()

BLUE = RGBColor(0x1F, 0x6F, 0xEB)  # same colour as \definecolor{principale} in the LaTeX report
INK = RGBColor(0x1E, 0x1E, 0x2E)
FONT = "Arial"


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def borders(table, color="B5C7F7"):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblpr.append(b)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)


def add_runs(par, text, size=None, bold=False, color=None):
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        is_code = part.startswith("`")
        is_bold = bold or (part.startswith("**") and part.endswith("**") and len(part) > 4)
        if is_code:
            part = part[1:-1]
        elif part.startswith("**") and part.endswith("**") and len(part) > 4:
            part = part[2:-2]
        part = part.replace("[[", "").replace("]]", "")
        r = par.add_run(part)
        r.bold = is_bold
        if is_code:
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x66)
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return par


def field(par, instr, placeholder=""):
    def el(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e
    r = par.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    r = par.add_run()
    t = el("w:instrText", **{"xml:space": "preserve"})
    t.text = f" {instr} "
    r._r.append(t)
    r = par.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    # placeholder: text (or list of lines) shown until the field is updated (LibreOffice keeps it as is)
    for n, line in enumerate([placeholder] if isinstance(placeholder, str) else placeholder):
        r = par.add_run()
        if n:
            r.add_break()
        r.add_text(line)
    r = par.add_run()
    r._r.append(el("w:fldChar", **{"w:fldCharType": "end"}))


def table(doc, rows):
    header, body = rows[0], rows[1:]
    n = len(header)
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    # column widths from content length (a column never narrower than its longest word)
    weights = []
    for i in range(n):
        col = [r[i] if i < len(r) else "" for r in rows]
        longest_word = max((len(w) for c in col for w in c.split()), default=4)
        weights.append(max(longest_word + 4, min(60, max(len(c) for c in col))))
    total = sum(weights)
    widths = [Cm(16.5 * w / total) for w in weights]
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips if hasattr(w, "twips") else w / 635)))
    for ri, row in enumerate([header] + body):
        cells = t.add_row().cells
        for ci in range(n):
            text = row[ci] if ci < len(row) else ""
            cell = cells[ci]
            cell.width = widths[ci]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, text, size=8.5, bold=ri == 0)
            if ri == 0:
                shade(cell, "EAEDF7")
        if ri == 0:
            repeat_header(t.rows[0])
    doc.add_paragraph()


def diagrams(doc):
    for png in sorted(PNG.glob("*.png")):
        num = png.name[:2]
        doc.add_picture(str(png), width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        default = png.stem[3:].replace("-", " ").capitalize() if num.isdigit() else png.stem
        add_runs(cap, f"Diagramme {num} : {CAPTIONS.get(num, default)}", size=9, color=INK).runs[0].italic = True


def setup(doc, title):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 26), ("Subtitle", 15), ("Heading 1", 16), ("Heading 2", 12.5)):
        s = doc.styles[name]
        s.font.name = FONT
        s.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(size)
        s.font.color.rgb = BLUE if name != "Subtitle" else INK
        s.font.bold = name != "Subtitle"
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2.25)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    hp = sec.header.paragraphs[0]
    add_runs(hp, f"{META.get('projet', 'Projet')} · {title} · v{META.get('version', '0.1')}", size=8,
             color=RGBColor(0x80, 0x80, 0x80))
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(fp, f"{META.get('pied', 'PFE ESPRIT')} · page ", size=8, color=RGBColor(0x80, 0x80, 0x80))
    field(fp, "PAGE", "1")
    # Word refreshes the table of contents when the file is opened
    s = doc.settings.element
    u = OxmlElement("w:updateFields")
    u.set(qn("w:val"), "true")
    s.append(u)


def convert(md):
    front = re.match(r"^---\n(.*?)\n---\n", md, flags=re.S)
    if front:
        pairs = (l.split(":", 1) for l in front.group(1).splitlines() if ":" in l)
        META.update({k.strip(): v.strip() for k, v in pairs})
    md = re.sub(r"^---\n.*?\n---\n", "", md, flags=re.S)
    lines = md.splitlines()
    sections = [l[3:] for l in lines if l.startswith("## ")][1:]
    doc = Document()
    setup(doc, "Cahier des charges")
    i, seen_subtitle, first_section = 0, False, True
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue
        if line.startswith("![["):
            pass  # Obsidian embed: the Word file shows the diagrams through <!-- diagrammes -->
        elif line.strip() == "<!-- diagrammes -->":
            diagrams(doc)
        elif line.startswith("# "):
            doc.add_paragraph()
            add_runs(doc.add_paragraph(style="Title"), line[2:])
        elif line.startswith("## ") and not seen_subtitle:
            add_runs(doc.add_paragraph(style="Subtitle"), line[3:])
            seen_subtitle = True
        elif line.startswith("## "):
            if first_section:
                doc.add_page_break()
                add_runs(doc.add_paragraph(), "Sommaire", size=16, bold=True, color=BLUE)  # not a heading: stays out of the TOC
                p = doc.add_paragraph()
                field(p, 'TOC \\o "1-2" \\h \\z \\u', sections)  # section list until Word updates the field
                doc.add_page_break()
                first_section = False
            add_runs(doc.add_heading(level=1), line[3:])
        elif line.startswith("### "):
            add_runs(doc.add_heading(level=2), line[4:])
        elif line.startswith("- [ ] ") or line.startswith("- [x] "):
            add_runs(doc.add_paragraph(style="List Bullet"), ("☑ " if "[x]" in line[:6] else "☐ ") + line[6:])
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1
    return doc


if __name__ == "__main__":
    document = convert(SRC.read_text(encoding="utf-8"))
    props = document.core_properties
    props.title = f"Cahier des charges : {META.get('projet', '')}"
    props.author = props.last_modified_by = META.get("auteur", "")
    document.save(OUT)
    print("written", OUT)
