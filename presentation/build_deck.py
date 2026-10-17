"""Construit presentation/Presentation.pptx (16:9) à partir de contenu.py, avec python-pptx.

Usage : python presentation/build_deck.py
PDF   : powershell -ExecutionPolicy Bypass -File scripts/export-pdf.ps1 presentation/Presentation.pptx
        (ou : soffice --headless --convert-to pdf presentation/Presentation.pptx)
Une image absente est remplacée par un cadre « Image à produire », comme dans le rapport LaTeX.
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from contenu import META, SLIDES  # noqa: E402

W, H = Emu(12192000), Emu(6858000)  # 13,33 x 7,5 pouces
MAIN = RGBColor.from_string(META.get("couleur", "1F6FEB"))
ACCENT = RGBColor.from_string(META.get("accent", "E8590C"))
INK, GREY, LIGHT, WHITE = (RGBColor.from_string(c) for c in ("1E1E2E", "6B7280", "F3F4F6", "FFFFFF"))
FONT = "Segoe UI"
M = Emu(560000)  # marge latérale


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    for i, line in enumerate(value if isinstance(value, list) else [value]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def picture(slide, rel, x, y, w, h):
    """Image centrée dans le cadre (x, y, w, h), proportions conservées."""
    path = ROOT / rel if rel else None
    if not path or not path.exists():
        r = rect(slide, x, y, w, h, LIGHT, line=ACCENT)
        text(slide, x, y, w, h, f"Image à produire : {rel}", 14, ACCENT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        return r
    pic = slide.shapes.add_picture(str(path), x, y)
    ratio = min(w / pic.width, h / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left, pic.top = int(x + (w - pic.width) / 2), int(y + (h - pic.height) / 2)
    return pic


def chrome(slide, n, title, subtitle=""):
    """En-tête (titre, sous-titre, filet) et pied de page communs."""
    rect(slide, M, Emu(470000), Emu(90000), Emu(560000), MAIN)
    text(slide, M + Emu(200000), Emu(380000), W - 2 * M, Emu(500000), title, 30, INK, True)
    if subtitle:
        text(slide, M + Emu(200000), Emu(880000), W - 2 * M, Emu(400000), subtitle, 16, GREY)
    text(slide, M, H - Emu(450000), Emu(8000000), Emu(300000), META.get("pied", ""), 10, GREY)
    text(slide, W - M - Emu(800000), H - Emu(450000), Emu(800000), Emu(300000), str(n), 10, GREY, align=PP_ALIGN.RIGHT)


BODY_Y = Emu(1500000)
BODY_H = H - BODY_Y - Emu(650000)


def s_titre(s, d, n):
    rect(s, 0, 0, W, H, MAIN)
    rect(s, M, Emu(2300000), Emu(120000), Emu(1700000), ACCENT)
    text(s, M + Emu(350000), Emu(2150000), W - 2 * M, Emu(1000000), d["title"], 54, WHITE, True)
    text(s, M + Emu(350000), Emu(3150000), W - 3 * M, Emu(900000), d.get("subtitle", ""), 22, WHITE)
    text(s, M + Emu(350000), H - Emu(1100000), W - 2 * M, Emu(400000), d.get("footer", ""), 14, WHITE)
    text(s, M + Emu(350000), H - Emu(700000), W - 2 * M, Emu(400000), META.get("pied", ""), 12, WHITE)


def s_plan(s, d, n):
    chrome(s, n, d["title"])
    items = d["items"]
    col = (len(items) + 1) // 2
    for i, item in enumerate(items):
        x = M + Emu(200000) + (i // col) * Emu(5600000)
        y = BODY_Y + Emu(200000) + (i % col) * Emu(1150000)
        text(s, x, y, Emu(1000000), Emu(800000), f"{i + 1:02d}", 36, MAIN, True)
        text(s, x + Emu(1100000), y + Emu(150000), Emu(4300000), Emu(800000), item, 20, INK)


def s_section(s, d, n):
    rect(s, 0, 0, W, H, LIGHT)
    text(s, M + Emu(200000), Emu(1700000), Emu(3000000), Emu(1500000), d.get("number", ""), 96, MAIN, True)
    text(s, M + Emu(200000), Emu(3400000), W - 2 * M, Emu(800000), d["title"], 40, INK, True)
    text(s, M + Emu(200000), Emu(4250000), W - 2 * M, Emu(600000), d.get("subtitle", ""), 20, GREY)
    text(s, W - M - Emu(800000), H - Emu(450000), Emu(800000), Emu(300000), str(n), 10, GREY, align=PP_ALIGN.RIGHT)


def s_puces(s, d, n):
    chrome(s, n, d["title"], d.get("subtitle", ""))
    has_img = bool(d.get("image"))
    w = Emu(5300000) if has_img else W - 2 * M - Emu(200000)
    for i, b in enumerate(d["bullets"]):
        y = BODY_Y + Emu(150000) + i * Emu(900000)
        rect(s, M + Emu(200000), y + Emu(150000), Emu(160000), Emu(160000), ACCENT, MSO_SHAPE.OVAL)
        text(s, M + Emu(550000), y, w - Emu(400000), Emu(800000), b, 20, INK)
    if has_img:
        picture(s, d["image"], M + Emu(6000000), BODY_Y, W - 2 * M - Emu(6000000), BODY_H)


def s_image(s, d, n):
    chrome(s, n, d["title"], d.get("subtitle", ""))
    picture(s, d["image"], M, BODY_Y, W - 2 * M, BODY_H - Emu(350000))
    if d.get("caption"):
        text(s, M, H - Emu(950000), W - 2 * M, Emu(300000), d["caption"], 12, GREY, align=PP_ALIGN.CENTER)


def s_cartes(s, d, n):
    chrome(s, n, d["title"], d.get("subtitle", ""))
    cards = d["cards"]
    per_row = 3 if len(cards) in (3, 5, 6) else 2
    rows = (len(cards) + per_row - 1) // per_row
    gap = Emu(250000)
    cw = int((W - 2 * M - (per_row - 1) * gap) / per_row)
    ch = int((BODY_H - (rows - 1) * gap) / rows)
    for i, (title, body) in enumerate(cards):
        x = M + (i % per_row) * (cw + gap)
        y = BODY_Y + (i // per_row) * (ch + gap)
        rect(s, x, y, cw, ch, LIGHT, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.06
        rect(s, x, y, cw, Emu(70000), MAIN)
        text(s, x + Emu(250000), y + Emu(300000), cw - Emu(500000), Emu(500000), title, 22, MAIN, True)
        text(s, x + Emu(250000), y + Emu(900000), cw - Emu(500000), ch - Emu(1000000), body, 16, INK)


def s_chiffres(s, d, n):
    chrome(s, n, d["title"], d.get("subtitle", ""))
    stats = d["stats"]
    cw = int((W - 2 * M) / len(stats))
    for i, (value, label) in enumerate(stats):
        x = M + i * cw
        text(s, x, BODY_Y + Emu(900000), cw, Emu(1100000), value, 60, MAIN if i % 2 == 0 else ACCENT, True, PP_ALIGN.CENTER)
        text(s, x + Emu(200000), BODY_Y + Emu(2100000), cw - Emu(400000), Emu(800000), label, 18, INK, align=PP_ALIGN.CENTER)


def s_frise(s, d, n):
    chrome(s, n, d["title"], d.get("subtitle", ""))
    steps = d["steps"]
    cw = int((W - 2 * M) / len(steps))
    line_y = BODY_Y + Emu(1500000)
    rect(s, M, line_y, W - 2 * M, Emu(50000), GREY)
    for i, (date, title, body) in enumerate(steps):
        cx = M + i * cw + cw // 2
        rect(s, cx - Emu(180000), line_y - Emu(155000), Emu(360000), Emu(360000), MAIN if i < len(steps) - 1 else ACCENT, MSO_SHAPE.OVAL)
        text(s, M + i * cw, line_y - Emu(1000000), cw, Emu(500000), date, 18, MAIN, True, PP_ALIGN.CENTER)
        text(s, M + i * cw + Emu(100000), line_y + Emu(450000), cw - Emu(200000), Emu(500000), title, 18, INK, True, PP_ALIGN.CENTER)
        text(s, M + i * cw + Emu(100000), line_y + Emu(950000), cw - Emu(200000), Emu(900000), body, 14, GREY, align=PP_ALIGN.CENTER)


def s_fin(s, d, n):
    rect(s, 0, 0, W, H, MAIN)
    text(s, M, Emu(2500000), W - 2 * M, Emu(1000000), d["title"], 44, WHITE, True, PP_ALIGN.CENTER)
    text(s, M, Emu(3600000), W - 2 * M, Emu(700000), d.get("subtitle", ""), 24, WHITE, align=PP_ALIGN.CENTER)


BUILDERS = {"titre": s_titre, "plan": s_plan, "section": s_section, "puces": s_puces, "image": s_image,
            "cartes": s_cartes, "chiffres": s_chiffres, "frise": s_frise, "fin": s_fin}


def build(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    for n, d in enumerate(SLIDES, 1):
        BUILDERS[d["type"]](prs.slides.add_slide(blank), d, n)
    prs.core_properties.title = META.get("projet", "")
    prs.core_properties.author = prs.core_properties.last_modified_by = META.get("etudiant", "")
    prs.save(out)
    print("written", out, f"({len(SLIDES)} slides)")


if __name__ == "__main__":
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Presentation.pptx")
